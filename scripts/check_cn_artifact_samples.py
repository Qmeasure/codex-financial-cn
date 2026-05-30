#!/usr/bin/env python3
"""Generate and inspect representative Chinese artifact samples."""
from __future__ import annotations

import sys
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "out"
OUT.mkdir(exist_ok=True)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
except Exception as exc:  # pragma: no cover
    print(f"requires openpyxl: {exc}", file=sys.stderr)
    sys.exit(2)

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
except Exception as exc:  # pragma: no cover
    print(f"requires python-pptx: {exc}", file=sys.stderr)
    sys.exit(2)


def fail(msg: str) -> None:
    print(f"FAIL: {msg}", file=sys.stderr)
    sys.exit(1)


xlsx = OUT / "cn_acceptance_sample.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "可比公司"
ws["A1"] = "A股优先可比公司分析"
ws["A1"].font = Font(name="Microsoft YaHei", bold=True)
ws["A3"] = "公司"
ws["B3"] = "收入（人民币亿元）"
ws["C3"] = "EV/EBITDA（倍）"
ws["D3"] = "来源"
for cell in ws[3]:
    cell.font = Font(name="Microsoft YaHei", bold=True)
ws["A4"] = "示例股份（600000.SH）"
ws["B4"] = 125.4
ws["C4"] = "=B4/10"
ws["D4"] = "交易所公告，2026年5月27日"
ws["D4"].hyperlink = "https://example.com/announcement"
for cell in ws[4]:
    cell.font = Font(name="Microsoft YaHei")
ws["B4"].number_format = '¥#,##0.0"亿"'
ws["C4"].number_format = '0.0x'
ws["A6"] = "免责声明：本样例不构成投资建议，应由人工审阅。"
ws["A6"].font = Font(name="Microsoft YaHei")
assumptions = wb.create_sheet("来源与假设")
assumptions["A1"] = "来源与假设"
assumptions["A1"].font = Font(name="Microsoft YaHei", bold=True)
assumptions["A2"] = "来源：交易所公告"
assumptions["A3"] = "假设：一致预期未作为正式输入"
checks = wb.create_sheet("检查区")
checks["A1"] = "检查项"
checks["B1"] = "状态"
checks["A2"] = "公式检查"
checks["B2"] = "通过"
wb.save(xlsx)

loaded = load_workbook(xlsx, data_only=False)
ws = loaded["可比公司"]
for sheet_name in ("可比公司", "来源与假设", "检查区"):
    if sheet_name not in loaded.sheetnames:
        fail(f"xlsx required sheet missing {sheet_name}")
if "人民币亿元" not in ws["B3"].value:
    fail("xlsx unit header missing Chinese RMB unit")
if ws["C4"].data_type != "f":
    fail("xlsx formula cell is not a formula")
if "投资建议" not in ws["A6"].value:
    fail("xlsx disclaimer missing")
if ws["A1"].font.name != "Microsoft YaHei" or ws["A4"].font.name not in (None, "Microsoft YaHei"):
    fail("xlsx Chinese font stack missing on title or body")
if not ws["D4"].hyperlink or "https://example.com/announcement" not in ws["D4"].hyperlink.target:
    fail("xlsx source hyperlink missing")

pptx = OUT / "cn_acceptance_sample.pptx"
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
title = slide.shapes.title
title.text = "A股优先估值摘要"
title.text_frame.paragraphs[0].runs[0].font.name = "Microsoft YaHei"
title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
box = slide.shapes.add_textbox(left=914400, top=1828800, width=7315200, height=1828800)
tf = box.text_frame
tf.text = "来源：交易所公告"
source_run = tf.paragraphs[0].runs[0]
source_run.font.name = "Microsoft YaHei"
source_run.font.size = Pt(6)
source_run.hyperlink.address = "https://example.com/announcement"
table_shape = slide.shapes.add_table(rows=2, cols=2, left=914400, top=3200400, width=3657600, height=731520)
table = table_shape.table
table.cell(0, 0).text = "指标"
table.cell(0, 1).text = "本期"
table.cell(1, 0).text = "收入（人民币亿元）"
table.cell(1, 1).text = "125.4"
table_source_box = slide.shapes.add_textbox(left=914400, top=3931920, width=3657600, height=228600)
table_source_tf = table_source_box.text_frame
table_source_tf.text = "来源：上市公司公告"
table_source_run = table_source_tf.paragraphs[0].runs[0]
table_source_run.font.name = "Microsoft YaHei"
table_source_run.font.size = Pt(6)
table_source_run.hyperlink.address = "https://example.com/announcement"
note_box = slide.shapes.add_textbox(left=914400, top=3566160, width=7315200, height=457200)
note_tf = note_box.text_frame
note_tf.text = "币种：人民币；单位：亿元；免责声明：不构成投资建议，应由人工审阅。"
note_run = note_tf.paragraphs[0].runs[0]
note_run.font.name = "Microsoft YaHei"
note_run.font.size = Pt(8)
chart_data = CategoryChartData()
chart_data.categories = ["1Q24", "2Q24", "3Q24"]
chart_data.add_series("收入", (87, 97, 99))
chart_shape = slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED,
    Inches(1.0),
    Inches(3.4),
    Inches(7.4),
    Inches(2.0),
    chart_data,
)
chart = chart_shape.chart
chart.has_legend = False
chart.value_axis.has_major_gridlines = False
if hasattr(chart.value_axis, "has_minor_gridlines"):
    chart.value_axis.has_minor_gridlines = False
prs.save(pptx)

prs2 = Presentation(pptx)
texts = "\n".join(shape.text for slide in prs2.slides for shape in slide.shapes if hasattr(shape, "text"))
if "A股优先估值摘要" not in texts:
    fail("pptx title missing Chinese text")
if "来源" not in texts or "人民币" not in texts:
    fail("pptx source or currency note missing")
if "免责声明" not in texts or "不构成投资建议" not in texts:
    fail("pptx disclaimer missing")
source_shapes = [
    shape for slide in prs2.slides for shape in slide.shapes
    if hasattr(shape, "text") and shape.text.startswith("来源：")
]
if len(source_shapes) != 2:
    fail("pptx exhibit source captions must cover chart and table samples")
for source_shape in source_shapes:
    source_text = source_shape.text
    if "；" in source_text or ";" in source_text or "PDD " in source_text:
        fail("pptx exhibit source caption must not list multiple sources")
    source_runs = source_shape.text_frame.paragraphs[0].runs
    if not source_runs or source_runs[0].font.size is None or source_runs[0].font.size.pt > 7:
        fail("pptx exhibit source caption must use very small font <= 7pt")
    if source_runs[0].hyperlink.address != "https://example.com/announcement":
        fail("pptx exhibit source caption hyperlink missing")
with zipfile.ZipFile(pptx) as zf:
    if "ppt/presentation.xml" not in zf.namelist():
        fail("pptx package missing presentation.xml")
    chart_parts = [name for name in zf.namelist() if name.startswith("ppt/charts/chart") and name.endswith(".xml")]
    if not chart_parts:
        fail("pptx chart sample missing chart XML")
    for chart_part in chart_parts:
        chart_xml = zf.read(chart_part).decode("utf-8")
        if "<c:majorGridlines" in chart_xml or "<c:minorGridlines" in chart_xml:
            fail(f"pptx bar chart gridlines must be disabled in {chart_part}")

docx = OUT / "cn_acceptance_sample.docx"
content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
  <Override PartName="/word/numbering.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"/>
</Types>
"""
package_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""
styles_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="paragraph" w:default="1" w:styleId="Normal">
    <w:name w:val="正文"/>
    <w:rPr>
      <w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/>
      <w:sz w:val="21"/>
    </w:rPr>
  </w:style>
  <w:style w:type="paragraph" w:styleId="Heading1">
    <w:name w:val="标题 1"/>
    <w:basedOn w:val="Normal"/>
    <w:rPr>
      <w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/>
      <w:b/>
      <w:color w:val="1F4E5F"/>
      <w:sz w:val="30"/>
    </w:rPr>
  </w:style>
</w:styles>
"""
numbering_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:abstractNum w:abstractNumId="1">
    <w:multiLevelType w:val="singleLevel"/>
    <w:lvl w:ilvl="0">
      <w:start w:val="1"/>
      <w:numFmt w:val="decimal"/>
      <w:lvlText w:val="%1."/>
      <w:lvlJc w:val="left"/>
      <w:pPr>
        <w:ind w:left="420" w:hanging="240"/>
      </w:pPr>
    </w:lvl>
  </w:abstractNum>
  <w:num w:numId="1">
    <w:abstractNumId w:val="1"/>
  </w:num>
</w:numbering>
"""
document_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdSource" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com/announcement" TargetMode="External"/>
</Relationships>
"""
document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <w:body>
    <w:p>
      <w:pPr><w:pStyle w:val="Heading1"/></w:pPr>
      <w:r>
        <w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr>
        <w:t>A股优先业绩更新报告</w:t>
      </w:r>
    </w:p>
    <w:p>
      <w:pPr>
        <w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>
      </w:pPr>
      <w:r>
        <w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr>
        <w:t>收入同比增长，毛利率应结合公告口径审阅。</w:t>
      </w:r>
    </w:p>
    <w:tbl>
      <w:tblPr>
        <w:tblW w:w="9000" w:type="dxa"/>
        <w:tblLayout w:type="fixed"/>
        <w:tblCellMar>
          <w:top w:w="80" w:type="dxa"/>
          <w:left w:w="120" w:type="dxa"/>
          <w:bottom w:w="80" w:type="dxa"/>
          <w:right w:w="120" w:type="dxa"/>
        </w:tblCellMar>
      </w:tblPr>
      <w:tblGrid>
        <w:gridCol w:w="3000"/>
        <w:gridCol w:w="3000"/>
        <w:gridCol w:w="3000"/>
      </w:tblGrid>
      <w:tr>
        <w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr><w:t>指标</w:t></w:r></w:p></w:tc>
        <w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr><w:t>本期</w:t></w:r></w:p></w:tc>
        <w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr><w:t>来源</w:t></w:r></w:p></w:tc>
      </w:tr>
      <w:tr>
        <w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr><w:t>收入（人民币亿元）</w:t></w:r></w:p></w:tc>
        <w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr><w:t>125.4</w:t></w:r></w:p></w:tc>
        <w:tc><w:tcPr><w:tcW w:w="3000" w:type="dxa"/></w:tcPr><w:p><w:r><w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr><w:t>数据来源：交易所公告</w:t></w:r></w:p></w:tc>
      </w:tr>
    </w:tbl>
    <w:p>
      <w:r>
        <w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/><w:sz w:val="12"/></w:rPr>
        <w:t>来源：</w:t>
      </w:r>
      <w:hyperlink r:id="rIdSource">
        <w:r>
          <w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/><w:sz w:val="12"/><w:u w:val="single"/><w:color w:val="0563C1"/></w:rPr>
          <w:t>交易所公告</w:t>
        </w:r>
      </w:hyperlink>
    </w:p>
    <w:p>
      <w:r>
        <w:rPr><w:rFonts w:eastAsia="Microsoft YaHei" w:ascii="Microsoft YaHei" w:hAnsi="Microsoft YaHei"/></w:rPr>
        <w:t>免责声明：本材料不构成投资建议，应由人工审阅。</w:t>
      </w:r>
    </w:p>
    <w:sectPr>
      <w:pgSz w:w="11906" w:h="16838"/>
      <w:pgMar w:top="1021" w:right="1021" w:bottom="907" w:left="1021" w:header="454" w:footer="454" w:gutter="0"/>
    </w:sectPr>
  </w:body>
</w:document>
"""

with zipfile.ZipFile(docx, "w", zipfile.ZIP_DEFLATED) as zf:
    zf.writestr("[Content_Types].xml", content_types)
    zf.writestr("_rels/.rels", package_rels)
    zf.writestr("word/document.xml", document_xml)
    zf.writestr("word/styles.xml", styles_xml)
    zf.writestr("word/numbering.xml", numbering_xml)
    zf.writestr("word/_rels/document.xml.rels", document_rels)

with zipfile.ZipFile(docx) as zf:
    names = set(zf.namelist())
    required_parts = {
        "word/document.xml",
        "word/styles.xml",
        "word/numbering.xml",
        "word/_rels/document.xml.rels",
    }
    missing = sorted(required_parts - names)
    if missing:
        fail(f"docx missing parts: {', '.join(missing)}")
    doc_xml = zf.read("word/document.xml").decode("utf-8")
    style_xml = zf.read("word/styles.xml").decode("utf-8")
    num_xml = zf.read("word/numbering.xml").decode("utf-8")
    rel_xml = zf.read("word/_rels/document.xml.rels").decode("utf-8")

for needle in (
    'w:eastAsia="Microsoft YaHei"',
    'w:ascii="Microsoft YaHei"',
    'w:hAnsi="Microsoft YaHei"',
):
    if needle not in doc_xml or needle not in style_xml:
        fail(f"docx font attribute missing {needle}")
for needle in ("<w:tblGrid>", "<w:tcW ", "<w:tblCellMar>"):
    if needle not in doc_xml:
        fail(f"docx table geometry missing {needle}")
if "<w:numPr>" not in doc_xml or "<w:abstractNum " not in num_xml or "<w:num " not in num_xml:
    fail("docx real numbering definitions missing")
doc_visible_text = "".join(ET.fromstring(doc_xml).itertext())
if "•" in doc_visible_text or "http://" in doc_visible_text or "https://" in doc_visible_text:
    fail("docx visible body text contains fake bullet or bare URL")
for forbidden in (
    "\u7ec8\u7aef\u8d1f\u8377",
    "\u7ec8\u7aef\u590d\u6838",
    "\u7ec8\u7aef\u5ba1\u6838",
    "\u9700\u7ec8\u7aef\u590d\u6838",
    "\u9700\u786e\u8ba4",
    "\u5f85\u786e\u8ba4",
    "\u590d\u6838",
):
    if forbidden in doc_visible_text:
        fail("docx visible text contains forbidden fixed review label")
if 'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"' not in rel_xml:
    fail("docx hyperlink relationship missing")
if 'Target="https://example.com/announcement"' not in rel_xml or 'TargetMode="External"' not in rel_xml:
    fail("docx hyperlink target missing or not external")
w_ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
doc_root = ET.fromstring(doc_xml)
source_paragraphs = []
for paragraph in doc_root.findall(".//w:p", w_ns):
    paragraph_text = "".join(paragraph.itertext()).strip()
    if paragraph_text.startswith("来源："):
        source_paragraphs.append(paragraph)
if len(source_paragraphs) != 1:
    fail("docx exhibit source caption must keep exactly one primary source")
source_text = "".join(source_paragraphs[0].itertext())
if "；" in source_text or ";" in source_text or "PDD " in source_text:
    fail("docx exhibit source caption must not list multiple sources")
hyperlinks = source_paragraphs[0].findall(".//w:hyperlink", w_ns)
if len(hyperlinks) != 1:
    fail("docx exhibit source caption must contain exactly one hyperlink")
for size in source_paragraphs[0].findall(".//w:sz", w_ns):
    val = size.attrib.get(f"{{{w_ns['w']}}}val")
    if val is None or int(val) > 14:
        fail("docx exhibit source caption font must be <= 7pt")
if '<w:sz w:val="12"/>' not in doc_xml:
    fail("docx source caption must default to 6pt")
for needle in ("交易所公告", "数据来源", "免责声明"):
    if needle not in doc_xml:
        fail(f"docx Chinese content missing {needle}")

md = OUT / "cn_acceptance_sample.md"
md.write_text(
    "# 中文研究摘要\n\n"
    "数据来源：[交易所公告](https://example.com/announcement)，2026年5月27日。\n\n"
    "币种与单位：人民币亿元。\n\n"
    "数据缺口：一致预期未作为正式结论。\n\n"
    "风险提示：本材料不构成投资建议，应由人工审阅。\n"
)
text = md.read_text()
for needle in ("数据来源", "人民币亿元", "数据缺口", "不构成投资建议"):
    if needle not in text:
        fail(f"markdown sample missing {needle}")
visible_markdown_text = re.sub(r"\[[^\]]+\]\(https?://[^)]+\)", "", text)
if "http://" in visible_markdown_text or "https://" in visible_markdown_text:
    fail("markdown sample contains bare visible URL")

print("OK — generated and inspected CN XLSX/PPTX/DOCX/Markdown artifact samples.")
